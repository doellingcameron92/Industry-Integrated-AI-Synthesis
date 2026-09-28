"""Scenario-based evaluation of the resume blinding service.

Usage:  python -m evaluation.run_resume_evaluation
Writes evaluation/resume_results.json and evaluation/resume_results.md; traces go to logs/resume_evaluation/.
Blinded documents approved during evaluation are written to a temporary directory, never to the repository.
"""

from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
from datetime import date
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from resume_blinder.parse import HeuristicExtractor  # noqa: E402
from resume_blinder.service import (  # noqa: E402
    BlindingRequest,
    BlindingResult,
    ResumeBlinderService,
    ServiceConfig,
    auto_approve,
    auto_deny,
)

EVAL_DIR = ROOT / "evaluation"
RESUME_DIR = EVAL_DIR / "resumes"
EVAL_LOG_DIR = "logs/resume_evaluation"
AS_OF = date(2026, 1, 1)


class HallucinatingRewriter:
    """Simulates an LLM rewriter that invents metrics and evaluative claims."""

    name = "HallucinatingRewriter"

    def rewrite(self, text: str) -> str:
        return f"{text.rstrip('.')}, increasing revenue by 500% as the strongest performer on the team."


class IdentityLeakingRewriter:
    """Simulates an LLM rewriter that re-inserts identity details taken from its context."""

    name = "IdentityLeakingRewriter"

    def __init__(self, leak: str) -> None:
        self.leak = leak

    def rewrite(self, text: str) -> str:
        return f"{text.rstrip('.')} ({self.leak})."


def write_docx_resume(text: str, path: Path, photo: bool = True) -> Path:
    """Build a DOCX resume from plain text: headshot + contact header, bullets as list paragraphs."""
    import docx
    from docx.shared import Inches
    from PIL import Image

    document = docx.Document()
    lines = text.splitlines()
    document.sections[0].header.paragraphs[0].text = lines[0]
    if photo:
        buf = io.BytesIO()
        Image.new("RGB", (32, 32), (180, 140, 120)).save(buf, format="PNG")
        buf.seek(0)
        document.add_picture(buf, width=Inches(0.8))
    for line in lines[1:]:
        if line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def write_pdf_resume(text: str, path: Path) -> Path:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(path), pagesize=LETTER, invariant=1)
    y = 750
    for line in text.splitlines():
        if y < 60:
            pdf.showPage()
            y = 750
        pdf.setFont("Helvetica", 9)
        pdf.drawString(50, y, line)
        y -= 13
    pdf.save()
    return path


def prepare_source(sc: dict[str, Any], workdir: Path) -> str:
    source = RESUME_DIR / f"{sc['resume']}.txt"
    fmt = sc.get("format", "txt")
    if fmt == "docx":
        return str(write_docx_resume(source.read_text(encoding="utf-8"), workdir / f"{sc['id']}.docx"))
    if fmt == "pdf":
        return str(write_pdf_resume(source.read_text(encoding="utf-8"), workdir / f"{sc['id']}.pdf"))
    return str(source)


def observe(result: BlindingResult) -> dict[str, Any]:
    trace = [json.loads(line) for line in Path(result.trace_path).read_text(encoding="utf-8").splitlines()]
    audit = result.draft.audit if result.draft else None
    return {
        "status": result.status,
        "refusal_reason": result.refusal_reason,
        "saved": bool(result.outputs),
        "confidence": result.confidence,
        "grounding_rate": audit.grounding_rate if audit else None,
        "ungrounded": len(audit.ungrounded) if audit else 0,
        "leakage_dropped": len(audit.leaked) if audit else 0,
        "document_leaks": len(audit.document_leaks) if audit else 0,
        "placeholder_artifacts": audit.placeholder_artifacts if audit else 0,
        "injection_flagged": any(r.get("kind") == "guardrail" and r.get("reason") == "prompt_injection_detected" for r in trace),
        "photos_removed": result.draft.removed.get("photo", 0) if result.draft else 0,
        "trace_records": len(trace),
    }


def evaluate_expectations(expect: dict[str, Any], observed: dict[str, Any], output: str) -> dict[str, bool]:
    checks: dict[str, bool] = {}
    for key, value in expect.items():
        if key == "max_confidence":
            checks[key] = observed["confidence"] <= value
        elif key == "min_confidence":
            checks[key] = observed["confidence"] >= value
        elif key.startswith("min_"):
            checks[key] = observed[key[4:]] >= value
        elif key == "output_contains":
            checks[key] = all(v.lower() in output.lower() for v in value)
        elif key == "output_not_contains":
            checks[key] = not any(v.lower() in output.lower() for v in value)
        else:
            checks[key] = observed.get(key) == value
    return checks


def run_scenarios() -> list[dict[str, Any]]:
    scenarios = json.loads((EVAL_DIR / "resume_scenarios.json").read_text(encoding="utf-8"))
    shutil.rmtree(ROOT / EVAL_LOG_DIR, ignore_errors=True)
    results = []
    with tempfile.TemporaryDirectory() as tmp:
        workdir = Path(tmp)
        config = ServiceConfig(log_dir=EVAL_LOG_DIR, output_dir=str(workdir / "approved"))
        for sc in scenarios:
            fault = sc.get("fault")
            rewriter = None
            if fault == "hallucinating_rewriter":
                rewriter = HallucinatingRewriter()
            elif fault == "identity_leaking_rewriter":
                rewriter = IdentityLeakingRewriter(sc["leak"])
            service = ResumeBlinderService(config=config, extractor=HeuristicExtractor(), rewriter=rewriter)
            request = BlindingRequest(source_path=prepare_source(sc, workdir), granularity=sc.get("granularity", "standard"),
                                      instruction=sc.get("instruction", "Produce a standardized blinded resume."), as_of=AS_OF)
            result = service.run(request, approval=auto_approve if sc.get("approval", "approve") == "approve" else auto_deny)
            output = result.draft.markdown if result.draft else result.message
            observed = observe(result)
            checks = evaluate_expectations(sc["expect"], observed, output)
            results.append({"id": sc["id"], "category": sc["category"], "resume": sc["resume"], "format": sc.get("format", "txt"),
                            "observed": observed, "checks": checks, "passed": all(checks.values()), "message": result.message})
            print(f"{'PASS' if all(checks.values()) else 'FAIL'}  {sc['id']}  conf={result.confidence:.2f}  status={result.status}")
    return results


def write_report(results: list[dict[str, Any]]) -> None:
    (EVAL_DIR / "resume_results.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    passed = sum(r["passed"] for r in results)
    lines = ["# Resume blinder evaluation results", "",
             f"Scenarios passed: **{passed}/{len(results)}**. Deterministic heuristic extractor; durations computed as of {AS_OF}.", "",
             "| ID | Category | Input | Status | Confidence | Grounding | Ungrounded dropped | Leak lines dropped | "
             "Document leaks | Saved | Result |",
             "|---|---|---|---|---|---|---|---|---|---|---|"]
    for r in results:
        o = r["observed"]
        grounding = "-" if o["grounding_rate"] is None else f"{o['grounding_rate']:.2f}"
        lines.append(f"| {r['id']} | {r['category']} | {r['resume']}.{r['format']} | {o['status']} | {o['confidence']:.2f} | "
                     f"{grounding} | {o['ungrounded']} | {o['leakage_dropped']} | {o['document_leaks']} | "
                     f"{'yes' if o['saved'] else 'no'} | {'PASS' if r['passed'] else 'FAIL'} |")
    failed = [(r["id"], [k for k, v in r["checks"].items() if not v]) for r in results if not r["passed"]]
    if failed:
        lines += ["", "## Failed checks", ""] + [f"- {sid}: {keys}" for sid, keys in failed]
    (EVAL_DIR / "resume_results.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    results = run_scenarios()
    write_report(results)
    print(f"\n{sum(r['passed'] for r in results)}/{len(results)} scenarios passed -> evaluation/resume_results.md")


if __name__ == "__main__":
    main()
