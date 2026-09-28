"""Regeneration: one uniform, skills-first template rendered to Markdown, then PDF/DOCX.

The template never contains placeholders for removed content: absent sections
render as ``None listed.`` and identity data has no slot in the template at all.
"""

from __future__ import annotations

import re
from pathlib import Path

from .schema import BlindedProfile

DOCUMENT_TITLE = "Standardized Resume"
DRAFT_BANNER = "> DRAFT — pending candidate approval. Do not save or share."
EMPTY_SECTION = "None listed."
SECTION_ORDER: tuple[tuple[str, str], ...] = (
    ("skills", "Skills"),
    ("experience", "Experience"),
    ("education", "Education"),
    ("certifications", "Certifications"),
    ("achievements", "Achievements"),
)
TEMPLATE_LINES = {f"# {DOCUMENT_TITLE}", DRAFT_BANNER, EMPTY_SECTION, *(f"## {title}" for _, title in SECTION_ORDER)}
FONT_CANDIDATES = ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/TTF/DejaVuSans.ttf")


def _clean(text: str) -> str:
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"^[#>\-*\s]+", "", text)
    return text.replace("*", "\\*").replace("_", "\\_")


def render_markdown(blinded: BlindedProfile, *, draft: bool = True) -> str:
    lines = [f"# {DOCUMENT_TITLE}", ""]
    if draft:
        lines += [DRAFT_BANNER, ""]
    for key, title in SECTION_ORDER:
        lines += [f"## {title}", ""]
        body: list[str] = []
        if key == "skills":
            if blinded.skills:
                body.append(" · ".join(_clean(s.text) for s in blinded.skills))
        elif key == "experience":
            for entry in blinded.experience:
                if entry.title is None:
                    continue
                body.append(f"### {_clean(entry.title.text)}")
                if entry.meta is not None:
                    body.append(f"*{_clean(entry.meta.text)}*")
                body += [f"- {_clean(b.text)}" for b in entry.bullets]
                body.append("")
        else:
            body += [f"- {_clean(line.text)}" for line in getattr(blinded, key)]
        while body and body[-1] == "":
            body.pop()
        lines += body or [EMPTY_SECTION]
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def _blocks(markdown: str) -> list[tuple[str, str]]:
    blocks = []
    for line in markdown.splitlines():
        if not line.strip():
            continue
        for prefix, kind in (("### ", "h3"), ("## ", "h2"), ("# ", "h1"), ("> ", "banner"), ("- ", "bullet")):
            if line.startswith(prefix):
                blocks.append((kind, line[len(prefix):]))
                break
        else:
            if line.startswith("*") and line.endswith("*") and len(line) > 1:
                blocks.append(("meta", line[1:-1]))
            else:
                blocks.append(("text", line))
    return [(kind, text.replace("\\*", "*").replace("\\_", "_")) for kind, text in blocks]


def render_pdf(markdown: str, path: str | Path) -> Path:
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    font = "Helvetica"
    for candidate in FONT_CANDIDATES:
        if Path(candidate).exists():
            if "ResumeSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ResumeSans", candidate))
            font = "ResumeSans"
            break
    styles = {
        "h1": ParagraphStyle("h1", fontName=font, fontSize=18, leading=22, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName=font, fontSize=13, leading=16, spaceBefore=10, spaceAfter=4),
        "h3": ParagraphStyle("h3", fontName=font, fontSize=11, leading=14, spaceBefore=6),
        "meta": ParagraphStyle("meta", fontName=font, fontSize=9, leading=12, textColor="#555555", spaceAfter=2),
        "text": ParagraphStyle("text", fontName=font, fontSize=10, leading=13),
        "banner": ParagraphStyle("banner", fontName=font, fontSize=9, leading=12, textColor="#aa0000", spaceAfter=6),
    }
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(out), pagesize=LETTER, leftMargin=0.8 * inch, rightMargin=0.8 * inch, topMargin=0.7 * inch,
                            bottomMargin=0.7 * inch, title=DOCUMENT_TITLE, author="", subject="", creator="resume_blinder",
                            invariant=1)
    story: list = []
    bullets: list = []

    def flush() -> None:
        if bullets:
            story.append(ListFlowable([ListItem(b, leftIndent=12) for b in bullets], bulletType="bullet", start="•",
                                      leftIndent=12, bulletFontName=font))
            bullets.clear()

    for kind, text in _blocks(markdown):
        safe = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if kind == "bullet":
            bullets.append(Paragraph(safe, styles["text"]))
            continue
        flush()
        story.append(Paragraph(safe, styles[kind]))
    flush()
    story.append(Spacer(1, 1))
    doc.build(story)
    return out


def render_docx(markdown: str, path: str | Path) -> Path:
    import docx
    from docx.shared import Pt, RGBColor

    document = docx.Document()
    props = document.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = DOCUMENT_TITLE
    props.comments = ""
    for kind, text in _blocks(markdown):
        if kind == "h1":
            document.add_heading(text, level=0)
        elif kind == "h2":
            document.add_heading(text, level=1)
        elif kind == "h3":
            document.add_heading(text, level=2)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        elif kind == "meta":
            run = document.add_paragraph().add_run(text)
            run.italic = True
            run.font.size = Pt(9)
        elif kind == "banner":
            run = document.add_paragraph().add_run(text)
            run.bold = True
            run.font.color.rgb = RGBColor(0xAA, 0, 0)
        else:
            document.add_paragraph(text)
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    return out
